import numpy as np
import pandas as pd
import os
import matplotlib.pyplot as plt
from sklearn.cross_decomposition import PLSRegression
from sklearn.model_selection import train_test_split
from modules.data_loader import load_data, load_metadata
from modules.preprocessing import Preprocessing
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

def sanitize_for_filename(s):
    import re
    s = str(s)
    s = re.sub(r'[\\/*?:"<>|]', '_', s)
    return s.replace(' ', '_')

def run_mccv_outlier_detection(
    train_data_dir,
    train_meta_path,
    test_data_dir,
    test_meta_path,
    filenames_per_column,
    output_dir=r"C:\Users\spect\Desktop\MD-Analysis-main (3)\MCCV results",
    target_columns=['target_SI', 'HGS_pp_avg', 'ADF_pp_avg'],
    n_iter=500,
    test_size=0.2,
    random_state=42
):
    os.makedirs(output_dir, exist_ok=True)
    master_outlier_log = []

    # ✅ Load spectra and filenames from both training and test sets
    _, _, X_train_raw, filenames_train = load_data(train_data_dir, return_filenames=True)
    _, _, X_test_raw, filenames_test = load_data(test_data_dir, return_filenames=True)
    
    # ✅ Combine spectra and filenames
    X_all_raw = np.hstack([X_train_raw, X_test_raw])
    filenames_per_column = filenames_train + filenames_test
    
    # ✅ Assert alignment between spectra and filenames
    assert X_all_raw.shape[1] == len(filenames_per_column), (
        f"Mismatch: X_all_raw has {X_all_raw.shape[1]} spectra, but {len(filenames_per_column)} filenames provided."
    )


    # ✅ Load and combine metadata
    y_train = load_metadata(train_meta_path)
    y_test = load_metadata(test_meta_path)
    y_all = pd.concat([y_train, y_test], ignore_index=True)
    
    # ✅ Assign target_SI = 0 to control samples
    y_all.loc[y_all['Type'] == 'Control', 'target_SI'] = 0
    
    # ✅ Extract sample IDs and source labels
    sample_ids = y_all['Sample_ID'].values
    source_labels = ['Train'] * len(y_train) + ['Test'] * len(y_test)


    preprocessing_combinations = [
        [],
        ['Normalization'],
        ['SNV'],
        ['EMSC'],
        ['Second Derivative'],
        ['Normalization', 'SNV'],
        ['Normalization', 'Second Derivative'],
        ['EMSC', 'Second Derivative'],
        ['SNV', 'Second Derivative'],
        ['Normalization', 'EMSC'],
        ['Normalization', 'EMSC', 'Second Derivative']
    ]

    for methods in preprocessing_combinations:
        preproc_name = '__'.join(methods) if methods else 'no_preprocessing'
        combo_output_dir = os.path.join(output_dir, sanitize_for_filename(preproc_name))
        os.makedirs(combo_output_dir, exist_ok=True)

        print(f"\n=== Running MCCV for preprocessing: {methods} ===")

        for col in target_columns:
            if col not in y_all.columns:
                print(f"Skipping {col}, not found.")
                continue

            sample_mask = ~y_all[col].isna()
            y_valid = y_all[col].values[sample_mask]
            ids_valid = sample_ids[sample_mask]
            src_valid = np.array(source_labels)[sample_mask]

            # Get which columns (spectra) correspond to valid samples
            valid_sample_ids = set(sample_ids[sample_mask])
            spectra_mask = [any(valid_id in fname for valid_id in valid_sample_ids) for fname in filenames_per_column]
            X_valid_raw = X_all_raw[:, spectra_mask].T


            n_samples = len(y_valid)
            spectra_per_sample = 9
            preds = np.full((n_samples, n_iter), np.nan)

            for i in range(n_iter):
                try:
                    sample_indices = np.arange(n_samples)
                    train_sample_idx, test_sample_idx = train_test_split(
                        sample_indices, test_size=test_size, random_state=random_state + i
                    )
                    train_spectrum_idx = np.hstack([
                        np.arange(s * spectra_per_sample, (s + 1) * spectra_per_sample) for s in train_sample_idx
                    ])
                    test_spectrum_idx = np.hstack([
                        np.arange(s * spectra_per_sample, (s + 1) * spectra_per_sample) for s in test_sample_idx
                    ])

                    X_train_raw_iter = X_valid_raw[train_spectrum_idx]
                    X_test_raw_iter = X_valid_raw[test_spectrum_idx]

                    X_train_proc = X_train_raw_iter.copy()
                    X_test_proc = X_test_raw_iter.copy()

                    if methods:
                        for method in methods:
                            if method == 'EMSC':
                                emsc_ref = np.mean(X_train_proc, axis=0)
                                X_train_proc = Preprocessing(None).emsc(X_train_proc.T, reference=emsc_ref).T
                                X_test_proc = Preprocessing(None).emsc(X_test_proc.T, reference=emsc_ref).T
                            else:
                                X_train_proc = Preprocessing(X_train_proc.T).preprocess([method]).T
                                X_test_proc = Preprocessing(X_test_proc.T).preprocess([method]).T

                    y_train_iter = np.repeat(y_valid[train_sample_idx], spectra_per_sample)

                    pls = PLSRegression(n_components=5)
                    pls.fit(X_train_proc, y_train_iter)
                    y_pred = pls.predict(X_test_proc).flatten()

                    for j, sample_idx in enumerate(test_sample_idx):
                        start = j * spectra_per_sample
                        end = (j + 1) * spectra_per_sample
                        pred_value = np.mean(y_pred[start:end])
                        preds[sample_idx, i] = pred_value
                    np.save(os.path.join(combo_output_dir, f"AllPredictions_{col}.npy"), preds)

                except Exception as e:
                    print(f"Iteration {i} failed: {e}")

            mean_pred = np.nanmean(preds, axis=1)
            std_pred = np.nanstd(preds, axis=1)
            abs_error = np.abs(y_valid - mean_pred)
            outlier_std = std_pred > 0.07
            outlier_error = abs_error > 0.3
            outlier_both = outlier_std & outlier_error

            # Get the true SampleIDs from the original y_all metadata using the same mask
            sample_ids_valid = y_all.loc[sample_mask, 'Sample_ID'].reset_index(drop=True)

            
            # Construct the DataFrame with correct SampleIDs
            df_log = pd.DataFrame({
                'SampleID': sample_ids_valid,
                'Source': src_valid,
                'True': y_valid,
                'Mean Prediction': mean_pred,
                'STD Prediction': std_pred,
                'Abs Error': abs_error,
                'Outlier_STD': outlier_std,
                'Outlier_Error': outlier_error,
                'Outlier_Both': outlier_both
            }).sort_values(by='Abs Error', ascending=False)


            df_top = df_log.head(10).copy()
            df_top['Target'] = col
            df_top['Preprocessing'] = ' + '.join(methods) if methods else 'None'
            master_outlier_log.append(df_top)

            # Save histogram
            plt.figure(figsize=(8, 6))
            plt.hist(abs_error, bins=20, color='gray', edgecolor='black')
            plt.axvline(0.3, color='red', linestyle='--', label='Error Threshold (0.3)')
            plt.xlabel("Absolute Prediction Error")
            plt.ylabel("Frequency")
            preproc_title = ' + '.join(methods) if methods else 'No Preprocessing'
            plt.title(f"Residual Histogram: {col} | {preproc_title}")

            plt.legend()
            plt.tight_layout()
            plt.savefig(os.path.join(combo_output_dir, f"Histogram_Abs_Error_{col}.png"), dpi=300)
            plt.show()
            plt.close()

            # Save MEAN-STD Plot
            plt.figure(figsize=(8, 6))
            for source in ['Train', 'Test']:
                mask = src_valid == source
                color = 'blue' if source == 'Train' else 'red'
                plt.scatter(mean_pred[mask], std_pred[mask], label=source, color=color, alpha=0.6, edgecolor='k')
            
            # Draw red dashed horizontal line at STD threshold
            plt.axhline(0.07, color='red', linestyle='--', linewidth=1.2, label='STD Threshold (0.07)')
            
            # Label only the outliers
            for i in range(n_samples):
                if outlier_both[i]:
                    plt.text(mean_pred[i], std_pred[i], str(ids_valid[i]), fontsize=6, alpha=0.8)
            
            plt.xlabel("Mean Prediction")
            plt.ylabel("Prediction STD")
            preproc_title = ' + '.join(methods) if methods else 'No Preprocessing'
            plt.title(f"MCCV MEAN-STD: {col} | {preproc_title}")

            plt.legend()
            plt.tight_layout()
            plt.savefig(os.path.join(combo_output_dir, f"MCCV_MEAN_STD_{col}.png"), dpi=300)
            plt.show()
            plt.close()
            
                        


            excel_path = os.path.join(combo_output_dir, f"MCCV_Outliers_{col}.xlsx")
            df_log.to_excel(excel_path, index=False)

            wb = load_workbook(excel_path)
            ws = wb.active
            std_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            error_fill = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
            both_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

            for row in range(2, ws.max_row + 1):
                if ws[f"K{row}"].value:
                    ws[f"K{row}"].fill = both_fill
                elif ws[f"I{row}"].value:
                    ws[f"I{row}"].fill = std_fill
                elif ws[f"J{row}"].value:
                    ws[f"J{row}"].fill = error_fill

            wb.save(excel_path)
            print(f"✓ Finished: {col} | Preproc: {methods} — Excel and plots saved.")

    # Save master outlier log
    if master_outlier_log:
        full_df = pd.concat(master_outlier_log, ignore_index=True)
        summary_path = os.path.join(output_dir, "Top_10_Outliers_Summary.xlsx")
        full_df.to_excel(summary_path, index=False)
        print(f"\n✓ Logged top outliers to summary file: {summary_path}")
